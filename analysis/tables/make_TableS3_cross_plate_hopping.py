#!/usr/bin/env python3
"""
Supplementary Table S3 -- cross-plate barcode hopping by plate of origin and read depth.

Supports the barcode-swapping paragraph of Results section 3 ("Design-aware decontamination
surgically removes ambient noise and rescues misclassified cells").

The verified totals are encoded in EXPECTED_TOTALS and the script hard-fails on drift, so a
silent data or definition change cannot ship into the manuscript unnoticed.

Definitions
-----------
Population   barcodes called single_clean or dirty_singlet in the with-design (WD) run of the
             maize and Arabidopsis library, carrying a plate-assigned expected_genome, with at
             least 100 pre-clean reads.
wrong top    top_genome differs from expected_genome.
pure wrong   wrong top AND expected_frac < 0.01, that is, under 1% of the barcode's reads sit
             on the genome its well was loaded with. This is the swapping-like class.

Why two files are joined
------------------------
cells_calls carries the call class. The composition table carries pre-clean total_reads,
top_genome, expected_genome and expected_frac. Do NOT substitute p_top2 from cells_calls for
expected_frac: it is computed on a different mass and gives 113 pure-wrong instead of 5.

Inputs (read-only, DATA = data/processed/scifiATAC_B73_Arabidopsis/SM2v2/decontam_with_design_alpha05_v2)
    SM2v2_cells_calls.decontam.tsv.gz
    SM2v2_pre_barcode_composition.tsv.gz

Writes TableS3 as .txt (TSV) and .xlsx into figures/tables/ by default.

Usage (from the repo root)
--------------------------
    python3 analysis/tables/make_TableS3_cross_plate_hopping.py [--outdir DIR]
"""

import argparse
import csv
import gzip
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
DATA = os.path.join(
    REPO, "data", "processed", "scifiATAC_B73_Arabidopsis", "SM2v2",
    "decontam_with_design_alpha05_v2",
)
CALLS_F = "SM2v2_cells_calls.decontam.tsv.gz"
COMP_F = "SM2v2_pre_barcode_composition.tsv.gz"

DEFAULT_OUTDIR = os.path.join(REPO, "figures", "tables")

SINGLET_CALLS = {"single_clean", "dirty_singlet"}
MIN_READS = 100
PURE_WRONG_MAX_EXPECTED_FRAC = 0.01

DEPTH_BINS = [
    (100, 500, "100 to 500"),
    (500, 1000, "500 to 1,000"),
    (1000, 5000, "1,000 to 5,000"),
    (5000, float("inf"), "5,000 and above"),
]

# Plate label as it should appear in the published table, keyed by expected_genome.
PLATE_LABEL = {"B73": "Maize", "At": "Arabidopsis"}
PLATE_ORDER = ["B73", "At"]

# Verified totals; the script self-checks against them so a silent data or definition drift
# cannot ship into the manuscript unnoticed.
EXPECTED_TOTALS = {
    "B73": {"n": 12445, "wrong": 188, "pure": 5},
    "At": {"n": 6709, "pure": 679},
}


def load_singlet_calls(path):
    calls = set()
    with gzip.open(path, "rt") as fh:
        for row in csv.DictReader(fh, delimiter="\t"):
            if row["call"] in SINGLET_CALLS:
                calls.add(row["barcode"])
    return calls


def tally(comp_path, singlets):
    """Return {(expected_genome, bin_label): [n, n_wrong, n_pure]}."""
    agg = {}
    with gzip.open(comp_path, "rt") as fh:
        for row in csv.DictReader(fh, delimiter="\t"):
            if row["barcode"] not in singlets:
                continue
            expected = row["expected_genome"].strip()
            if not expected:
                continue
            total = float(row["total_reads"])
            if total < MIN_READS:
                continue
            wrong = row["top_genome"].strip() != expected
            pure = wrong and float(row["expected_frac"]) < PURE_WRONG_MAX_EXPECTED_FRAC
            for lo, hi, label in DEPTH_BINS:
                if lo <= total < hi:
                    cell = agg.setdefault((expected, label), [0, 0, 0])
                    cell[0] += 1
                    cell[1] += wrong
                    cell[2] += pure
                    break
    return agg


def build_rows(agg):
    rows = []
    totals = {}
    for genome in PLATE_ORDER:
        running = [0, 0, 0]
        for _, _, label in DEPTH_BINS:
            n, wrong, pure = agg.get((genome, label), [0, 0, 0])
            for i, v in enumerate((n, wrong, pure)):
                running[i] += v
            rows.append([
                PLATE_LABEL[genome], label, n, wrong, pure,
                round(100.0 * pure / n, 2) if n else "",
            ])
        rows.append([
            PLATE_LABEL[genome], "All depths", running[0], running[1], running[2],
            round(100.0 * running[2] / running[0], 2) if running[0] else "",
        ])
        totals[genome] = running
    return rows, totals


def check(totals):
    """Fail loudly rather than ship a table that drifted from the verified values."""
    problems = []
    for genome, want in EXPECTED_TOTALS.items():
        n, wrong, pure = totals[genome]
        got = {"n": n, "wrong": wrong, "pure": pure}
        for key, value in want.items():
            if got[key] != value:
                problems.append(
                    f"{PLATE_LABEL[genome]} {key}: got {got[key]:,}, expected {value:,}"
                )
    if problems:
        sys.exit("REGRESSION against the verified values:\n  " + "\n  ".join(problems))


HEADER = [
    "Plate of origin", "Pre-clean read depth", "Singlet barcodes",
    "Wrong top genome", "Pure wrong", "Pure wrong (%)",
]

CAPTION = (
    "Table S3. Cross-plate barcode hopping by plate of origin and read depth. "
    "Population: barcodes from the with-design (WD) run of the maize and Arabidopsis "
    "library called single_clean or "
    "dirty_singlet, carrying a plate-assigned expected genome, with at least 100 pre-clean reads. "
    "\"Wrong top genome\" counts barcodes whose most-supported genome differs from the genome "
    "their well was loaded with. \"Pure wrong\" additionally requires under 1% of the barcode's "
    "reads on the loaded genome, the pattern expected of a swapped barcode. The maize-plate "
    "direction is the one that can be measured without confounding, because little Arabidopsis "
    "chromatin is available to fill a maize-plate barcode by ambient background. The "
    "Arabidopsis-plate rate falls with depth, which is the signature of ambient filling rather "
    "than of a barcode-level labelling error, since a labelling error would be depth-independent."
)


def write_txt(path, rows):
    with open(path, "w", newline="") as fh:
        writer = csv.writer(fh, delimiter="\t", lineterminator="\n")
        writer.writerow(HEADER)
        writer.writerows(rows)
        fh.write("\n" + CAPTION + "\n")


def write_xlsx(path, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        print("openpyxl not available, skipping .xlsx", file=sys.stderr)
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "Table S3"
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        if str(row[1].value) == "All depths":
            for cell in row:
                cell.font = Font(bold=True)
        row[5].number_format = "0.00"
    for column, width in zip("ABCDEF", (16, 18, 17, 18, 12, 15)):
        ws.column_dimensions[column].width = width

    caption_row = ws.max_row + 2
    ws.cell(row=caption_row, column=1, value=CAPTION).alignment = Alignment(
        wrap_text=True, vertical="top"
    )
    ws.merge_cells(start_row=caption_row, start_column=1, end_row=caption_row + 6, end_column=6)
    wb.save(path)
    return True


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--outdir", default=DEFAULT_OUTDIR)
    args = parser.parse_args()

    calls_path = os.path.join(DATA, CALLS_F)
    comp_path = os.path.join(DATA, COMP_F)
    for path in (calls_path, comp_path):
        if not os.path.exists(path):
            sys.exit(f"missing input: {path}")

    singlets = load_singlet_calls(calls_path)
    print(f"single_clean + dirty_singlet barcodes: {len(singlets):,}")

    rows, totals = build_rows(tally(comp_path, singlets))
    check(totals)
    print("self-check against the verified values: PASS")

    os.makedirs(args.outdir, exist_ok=True)
    txt = os.path.join(args.outdir, "TableS3_cross_plate_barcode_hopping.txt")
    xlsx = os.path.join(args.outdir, "TableS3_cross_plate_barcode_hopping.xlsx")
    write_txt(txt, rows)
    print(f"wrote {txt}")
    if write_xlsx(xlsx, rows):
        print(f"wrote {xlsx}")

    width = max(len(str(r[1])) for r in rows)
    for row in rows:
        pct = f"{row[5]:.2f}%" if row[5] != "" else "-"
        print(f"{row[0]:>12}  {row[1]:<{width}}  n={row[2]:>6,}  wrong={row[3]:>5,}  "
              f"pure={row[4]:>5,}  {pct:>7}")


if __name__ == "__main__":
    main()
