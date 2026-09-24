#!/usr/bin/env python3
"""
Supplementary Table S2 -- dataset inventory.

Cited twice by the main Methods, "scifi-ATAC-seq datasets".

What the table answers
----------------------
One row per dataset analysed in the paper: what it is, which reference genomes it was mapped
against, whether AmbientMapper was run with or without the plate design, what the dataset is
there to test, where it came from, and which figures use it.

Everything that can be read from a configuration file is read from one. The reference panels
come from the AmbientMapper run configs, the accessions from the SRA metadata sheet, the well
counts from the plate-design files, and the cleaning modes from the run directories that exist
on disk. Only the prose columns (material, role, figures) are curated here.

Datasets deliberately absent
----------------------------
A maize seedling library and a B73/Oh43 library from early drafts appear in no figure and have
no AmbientMapper run in this study, so neither has a row.

Inputs (all read-only)
----------------------
    config/{SM2v2,Root1_rep1,B73Mo17_rep1,B73Mo17_rep2,multiGenotypes_rep1}.ambientmapper.json
    config/scifi_Metadata_sra.clean.txt     (TSV, columns SampleID, Run, LibraryLayout)
    config/srr_root.txt                     (every maize-root run the project fetched, one per line)
    config/PlateDesign_SM2_ATAC.txt, config/PlateDesign_scifi_{B73Mo17_rep1,B73Mo17_rep2,multi_genotypes}.txt
    data/processed/<dataset>/decontam_*/    (directory existence decides the cleaning-mode column)
    data/processed/synthetic/{synthetic,synthetic_disc}/   (alpha_* datasets, barcoded/templates.tsv)
Optional: --root-fastq <demultiplexed R1 FASTQ of the maize root library>, see verify_root_fastq().

Writes TableS2 as .txt (TSV) and .xlsx into figures/tables/ by default.

Usage (from the repo root)
--------------------------
    python3 analysis/tables/make_TableS2_dataset_inventory.py [--outdir DIR] [--root-fastq FASTQ]
"""

import argparse
import csv
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
CONFIG_DIR = os.path.join(REPO, "config")
DATA = os.path.join(REPO, "data", "processed")
DEFAULT_OUTDIR = os.path.join(REPO, "figures", "tables")

AM_CONFIGS = CONFIG_DIR
SRA_SHEET = os.path.join(CONFIG_DIR, "scifi_Metadata_sra.clean.txt")
ROOT_SRR_FILE = os.path.join(CONFIG_DIR, "srr_root.txt")
SYNTH_DIR = os.path.join(DATA, "synthetic", "synthetic")
SYNTH_DISC_DIR = os.path.join(DATA, "synthetic", "synthetic_disc")

# Where each dataset's processed AmbientMapper outputs live under DATA.
DATASET_DIR = {
    "SM2": os.path.join(DATA, "scifiATAC_B73_Arabidopsis", "SM2v2"),
    "Root1": os.path.join(DATA, "marand2021_B73_root", "Root1_rep1"),
    "B73Mo17_rep1": os.path.join(DATA, "zhang2024", "B73Mo17_rep1"),
    "B73Mo17_rep2": os.path.join(DATA, "zhang2024", "B73Mo17_rep2"),
    "multiGenotypes_rep1": os.path.join(DATA, "zhang2024", "multiGenotypes_rep1"),
}

# GEO series of the maize and Arabidopsis library generated in this study.
SM2_ACCESSION = "GSE348516"

# Verified against NCBI E-utilities, not transcribed from a reference manager.
#   38589969  Zhang X et al. 2024, Genome Biol 25(1):90, scifi-ATAC-seq
#   33964211  Marand AP et al. 2021, Cell 184(11):3041-3055.e21, maize cis-regulatory atlas
PUBMED = {"zhang2024": "38589969", "marand2021": "33964211"}
BIOPROJECT = {"zhang2024": "PRJNA996051", "marand2021": "PRJNA648930"}

# The maize root library is GEO sample GSM4696884, titled "Root1", under PRJNA648930. That sample
# has two sequencing runs and ONLY THE FIRST was mapped and analysed here. Provenance:
#   - the demultiplexed reads carry the accession in the read name: the Root1_rep1 R1 FASTQ
#     begins "@SRR12331466.1_..." and the Root1_rep2 R1 FASTQ begins "@SRR12331467.1_..."
#   - the BAM header shows bwa mem was given the Root1_rep1 reads only; Root1_rep2 was never mapped
#   - srr_root.txt additionally lists SRR12331468 and SRR12331469, which esummary resolves to
#     GSM4696885 ("Root2"), a different sample that is not part of this study
ROOT_GEO_SAMPLE = "GSM4696884"
ROOT_RUNS = ["SRR12331466"]
ROOT_RUN_UNUSED = "SRR12331467"          # second run of the same GEO sample, not analysed here
ROOT_RUNS_EXCLUDED = ["SRR12331468", "SRR12331469"]   # GSM4696885, a different sample

# Where each dataset's reference panel is declared.
PANEL_CONFIG = {
    "SM2": "SM2v2.ambientmapper.json",
    "Root1": "Root1_rep1.ambientmapper.json",
    "B73Mo17_rep1": "B73Mo17_rep1.ambientmapper.json",
    "B73Mo17_rep2": "B73Mo17_rep2.ambientmapper.json",
    "multiGenotypes_rep1": "multiGenotypes_rep1.ambientmapper.json",
}

# Decontamination run directories (relative to DATASET_DIR), checked for existence so the mode
# column cannot go stale.
DECONTAM_DIRS = {
    "SM2": {
        "WD": "decontam_with_design_alpha05_v2",
        "ND": "decontam_without_design_alpha05_v2",
    },
    "Root1": {},
    "B73Mo17_rep1": {"ND": "decontam_without_design_alpha05_C0"},
    "B73Mo17_rep2": {"ND": "decontam_without_design_alpha05_C0"},
    "multiGenotypes_rep1": {"ND": "decontam_without_design_alpha05_C0"},
}

# Plate-design files, used only for the well count. Root1 has no combinatorial plate.
WELL_SOURCE = {
    "SM2": ("range", os.path.join(CONFIG_DIR, "PlateDesign_SM2_ATAC.txt")),
    "B73Mo17_rep1": ("perwell", os.path.join(CONFIG_DIR, "PlateDesign_scifi_B73Mo17_rep1.txt")),
    "B73Mo17_rep2": ("perwell", os.path.join(CONFIG_DIR, "PlateDesign_scifi_B73Mo17_rep2.txt")),
    "multiGenotypes_rep1": (
        "perwell",
        os.path.join(CONFIG_DIR, "PlateDesign_scifi_multi_genotypes.txt"),
    ),
}

# Curated prose. Kept here so every editable string is in one place.
PROSE = {
    "SM2": {
        "label": "Maize and Arabidopsis",
        "pmid": "Not applicable",
        "assay": "scifi-ATAC-seq",
        "material": (
            "Zea mays B73, V2 stage, above-ground organs, and Arabidopsis thaliana Columbia, "
            "whole seedling including root"
        ),
        "role": (
            "Interspecies ground truth. Contamination is directly measurable because the two "
            "genomes are unambiguously distinguishable, and the plate index records which "
            "species each well was loaded with. Used to calibrate decision thresholds, to "
            "benchmark decontamination, and to measure the biological effect of cleaning"
        ),
        "source": "This study",
        "accession": SM2_ACCESSION,
        "figures": "1, 2, 3, 5, S1, S2, S3, S4, S7, S8",
    },
    "Root1": {
        "label": "Maize root",
        "pmid": PUBMED["marand2021"],
        "assay": "scATAC-seq (10x)",
        "material": "Zea mays B73 root",
        "role": (
            "Zero-contamination control and extreme reference-redundancy stress test. The "
            "library is a single genotype mapped against 26 NAM founder genomes at roughly 95% "
            "sequence identity, so every non-B73 call is an error by construction"
        ),
        "source": "Marand et al. 2021",
        "accession": None,
        "figures": "4D to 4G, S6",
    },
    "B73Mo17_rep1": {
        "label": "B73/Mo17, replicate 1",
        "pmid": PUBMED["zhang2024"],
        "assay": "scifi-ATAC-seq",
        "material": "Zea mays B73 and Mo17, pooled in a single nuclei extraction",
        "role": (
            "Same-species generalizability. No plate index separates the two genotypes, so "
            "assignment is design-free. Used for the comparison against Souporcell and for the "
            "allele-resolution measurement of cleaning on WASP-corrected alignments"
        ),
        "source": "Zhang et al. 2024",
        "accession": None,
        "figures": "4H, 4J, 4L to 4N",
    },
    "B73Mo17_rep2": {
        "label": "B73/Mo17, replicate 2",
        "pmid": PUBMED["zhang2024"],
        "assay": "scifi-ATAC-seq",
        "material": "Zea mays B73 and Mo17, pooled in a single nuclei extraction",
        "role": "Independent replicate of the library above, analysed identically",
        "source": "Zhang et al. 2024",
        "accession": None,
        "figures": "4H, 4J, 4L to 4N",
    },
    "multiGenotypes_rep1": {
        "label": "Multi-genotype",
        "pmid": PUBMED["zhang2024"],
        "assay": "scifi-ATAC-seq",
        "material": "Seven maize NAM genotypes, each occupying its own set of plate wells",
        "role": (
            "Multi-genotype scalability. AmbientMapper was run design-free and the plate "
            "assignment was held back as independent ground truth against which the resulting "
            "genotype calls were scored"
        ),
        "source": "Zhang et al. 2024",
        "accession": None,
        "figures": "4I, 4K, 4L to 4N",
    },
    "synthetic": {
        "label": "Synthetic benchmark, all ortholog peaks",
        "pmid": "Not applicable",
        "assay": "Simulated 75 bp paired-end reads",
        "material": (
            "Reads simulated with ART from ortholog peak triplets defined on accessible regions "
            "called in the maize root library alignments"
        ),
        "role": (
            "Fully traceable titration. The genome of origin of every read is known, so "
            "sensitivity and precision can be measured against exact truth across a "
            "contamination series"
        ),
        "source": "This study, simulated",
        "accession": "Not applicable",
        "figures": "4A to 4C, S5",
    },
    "synthetic_disc": {
        "label": "Synthetic benchmark, discriminative peaks",
        "pmid": "Not applicable",
        "assay": "Simulated 75 bp paired-end reads",
        "material": (
            "As above, restricted to peaks carrying at least one discriminating variant per "
            "75 bp read"
        ),
        "role": (
            "Companion titration that isolates model behaviour from reference informativeness, "
            "by removing regions where no read can distinguish the candidate genomes"
        ),
        "source": "This study, simulated",
        "accession": "Not applicable",
        "figures": "4A to 4C, S5",
    },
}

ROW_ORDER = [
    "SM2", "Root1", "B73Mo17_rep1", "B73Mo17_rep2", "multiGenotypes_rep1",
    "synthetic", "synthetic_disc",
]

# Verified structure. Self-checked so a silent config change cannot ship.
EXPECTED = {
    "SM2": {"genomes": 2, "wells": 64},
    "Root1": {"genomes": 26, "wells": None},
    "B73Mo17_rep1": {"genomes": 2, "wells": 96},
    "B73Mo17_rep2": {"genomes": 2, "wells": 96},
    "multiGenotypes_rep1": {"genomes": 7, "wells": 96},
}
EXPECTED_SRR = {
    "B73Mo17_rep1": ["PRJNA996051", "SRR25320545", "SRR25320546", "SRR25320547"],
    "B73Mo17_rep2": ["PRJNA996051", "SRR25320539", "SRR25320540", "SRR25320541"],
    "multiGenotypes_rep1": ["PRJNA996051", "SRR25320542", "SRR25320543", "SRR25320544"],
}
EXPECTED_SYNTH = {"datasets": 15, "templates": 208, "genomes": 3}

WELL_RE = re.compile(r"^([A-H])(\d+)(?:-(\d+))?$")


def expand_wells(spec):
    wells = []
    for token in spec.split(","):
        token = token.strip()
        if not token:
            continue
        m = WELL_RE.match(token)
        if not m:
            raise ValueError(f"unparseable well token: {token!r}")
        row, start, end = m.group(1), int(m.group(2)), m.group(3)
        end = int(end) if end else start
        wells.extend(f"{row}{i}" for i in range(start, end + 1))
    return wells


def count_wells(kind, path):
    if not os.path.exists(path):
        sys.exit(f"missing input: {path}")
    wells = set()
    with open(path) as fh:
        for line in fh:
            line = line.rstrip("\n")
            if not line.strip():
                continue
            left, right = line.split("\t")
            wells.update(expand_wells(right) if kind == "range" else [right.strip()])
    return len(wells)


def read_panel(filename):
    path = os.path.join(AM_CONFIGS, filename)
    if not os.path.exists(path):
        sys.exit(f"missing input: {path}")
    with open(path) as fh:
        genomes = list(json.load(fh).get("genomes", {}).keys())
    if not genomes:
        sys.exit(f"{path} declares no genomes")
    return sorted(genomes, key=lambda g: (g != "B73", g.lower()))


def read_root_runs(root_fastq):
    """Verify the maize-root runs against config/srr_root.txt.

    That file holds every root run the project ever fetched, which spans two GEO samples. The
    table cites only GSM4696884 ("Root1"), so the selection is deliberate and is checked rather
    than taken wholesale.
    """
    if not os.path.exists(ROOT_SRR_FILE):
        sys.exit(f"missing input: {ROOT_SRR_FILE}")
    with open(ROOT_SRR_FILE) as fh:
        listed = {line.strip() for line in fh if line.strip()}
    missing = [r for r in ROOT_RUNS if r not in listed]
    if missing:
        sys.exit(
            f"{ROOT_SRR_FILE} no longer lists {', '.join(missing)}.\n"
            "  Table S2 cites GSM4696884 (Root1). Re-verify the accessions before shipping."
        )
    extra = sorted(listed - set(ROOT_RUNS) - {ROOT_RUN_UNUSED} - set(ROOT_RUNS_EXCLUDED))
    if extra:
        sys.exit(
            f"{ROOT_SRR_FILE} lists unrecognised runs: {', '.join(extra)}.\n"
            "  Resolve which GEO sample they belong to before shipping the table."
        )
    verify_root_fastq(root_fastq)
    return [BIOPROJECT["marand2021"], ROOT_GEO_SAMPLE] + ROOT_RUNS


def verify_root_fastq(path):
    """Optional check of the cited run against the read headers of the FASTQ that was mapped.

    The demultiplexed reads keep their SRA accession in the read name, so this turns the
    accession from an inference into a check. Only the first gzip block is read. The step runs
    only when --root-fastq points at an existing file: the FASTQ is about 100 GB and is not part
    of this repository, and an absent file is not evidence of a wrong accession. A mismatch on a
    readable file is a hard failure.
    """
    cited = ROOT_RUNS[0]
    if path is None:
        print(
            f"NOTE: root FASTQ header check skipped (no --root-fastq given). The cited run "
            f"{cited} rests on the provenance recorded at the top of this script.",
            file=sys.stderr,
        )
        return
    if not os.path.exists(path):
        print(
            f"NOTE: root FASTQ header check skipped, {path} does not exist. The cited run "
            f"{cited} rests on the provenance recorded at the top of this script.",
            file=sys.stderr,
        )
        return
    import gzip
    try:
        with gzip.open(path, "rt") as fh:
            header = fh.readline().strip()
    except OSError as exc:
        print(f"NOTE: could not read {path} ({exc}), root FASTQ header check skipped",
              file=sys.stderr)
        return
    if not header.lstrip("@").startswith(cited + "."):
        sys.exit(
            f"root accession mismatch.\n"
            f"  Table S2 cites {cited}, but {path} begins:\n    {header[:120]}\n"
            "  Fix ROOT_RUNS before shipping."
        )
    print(f"root accession {cited} verified against the read headers of {path}")


def read_accessions():
    if not os.path.exists(SRA_SHEET):
        sys.exit(f"missing input: {SRA_SHEET}")
    by_dataset = {}
    with open(SRA_SHEET) as fh:
        for row in csv.DictReader(fh, delimiter="\t"):
            sample = row["SampleID"].strip()
            m = re.match(r"^scifi_(.+)_\d+$", sample)
            if not m:
                continue
            by_dataset.setdefault(m.group(1), set()).add(row["Run"].strip())
    return {
        k: [BIOPROJECT["zhang2024"]] + sorted(v) for k, v in by_dataset.items()
    }


def _scan_synthetic_tree(root):
    if not os.path.isdir(root):
        sys.exit(f"missing input: {root}")
    datasets = sorted(d for d in os.listdir(root) if d.startswith("alpha_"))
    alphas = sorted({int(d.split("_")[1]) for d in datasets})
    templates_path = os.path.join(root, "barcoded", "templates.tsv")
    n_templates = None
    if os.path.exists(templates_path):
        with open(templates_path) as fh:
            n_templates = sum(1 for _ in fh) - 1
    return len(datasets), n_templates, alphas


def read_synthetic():
    """Return (n_titration_datasets, n_templates, alphas) for the synthetic benchmark.

    Both tracks share one titration design, so the two trees are scanned and compared. The
    table states a single structure for both rows, and that is only honest if they agree.
    """
    full = _scan_synthetic_tree(SYNTH_DIR)
    disc = _scan_synthetic_tree(SYNTH_DISC_DIR)
    if full != disc:
        sys.exit(
            "the two synthetic tracks no longer share a titration design:\n"
            f"  synthetic:      {full[0]} datasets, {full[1]} templates, alphas {full[2]}\n"
            f"  synthetic_disc: {disc[0]} datasets, {disc[1]} templates, alphas {disc[2]}\n"
            "  Table S2 states one structure for both rows and must be split first."
        )
    return full


def modes_on_disk(dataset):
    found = []
    for mode, rel in DECONTAM_DIRS[dataset].items():
        if os.path.isdir(os.path.join(DATASET_DIR[dataset], rel)):
            found.append(mode)
    return sorted(found, reverse=True)  # WD before ND


def check(panels, wells, accessions, synth):
    problems = []
    for dataset, want in EXPECTED.items():
        if len(panels[dataset]) != want["genomes"]:
            problems.append(
                f"{dataset} reference genomes: got {len(panels[dataset])}, "
                f"expected {want['genomes']}"
            )
        if want["wells"] is not None and wells.get(dataset) != want["wells"]:
            problems.append(
                f"{dataset} wells: got {wells.get(dataset)}, expected {want['wells']}"
            )
    for dataset, want in EXPECTED_SRR.items():
        got = accessions.get(dataset, [])
        if got != want:
            problems.append(f"{dataset} accessions: got {got}, expected {want}")

    n_datasets, n_templates, _ = synth
    if n_datasets != EXPECTED_SYNTH["datasets"]:
        problems.append(
            f"synthetic titration datasets: got {n_datasets}, "
            f"expected {EXPECTED_SYNTH['datasets']}"
        )
    if n_templates is not None and n_templates != EXPECTED_SYNTH["templates"]:
        problems.append(
            f"synthetic template barcodes: got {n_templates}, "
            f"expected {EXPECTED_SYNTH['templates']}"
        )

    if modes_on_disk("SM2") != ["WD", "ND"]:
        problems.append(f"SM2 cleaning modes on disk: {modes_on_disk('SM2')}, expected WD and ND")
    for dataset in ("B73Mo17_rep1", "B73Mo17_rep2", "multiGenotypes_rep1"):
        if modes_on_disk(dataset) != ["ND"]:
            problems.append(
                f"{dataset} cleaning modes on disk: {modes_on_disk(dataset)}, expected ND only"
            )

    if problems:
        sys.exit("REGRESSION against the verified dataset inventory:\n  " + "\n  ".join(problems))


HEADER = [
    "Dataset", "Assay", "Material", "Reference genomes (n)", "Reference panel",
    "Combinatorial Tn5 wells", "AmbientMapper mode", "Role in this study",
    "Source", "PubMed ID", "Data accession", "Figures",
]

CAPTION = (
    "Table S2. Inventory of the datasets analysed in this study. Reference panel lists the "
    "genomes each library was mapped against independently, which is the candidate set "
    "AmbientMapper resolves each barcode over. AmbientMapper mode records whether the plate "
    "design was supplied to the run as a prior, WD, or withheld so that genotype was inferred "
    "from the data alone, ND. Only the maize and Arabidopsis library was run in both modes, and "
    "that pair is the WD versus ND comparison used throughout the paper. The multi-genotype "
    "library has a usable plate design but was run design-free, with the design held back as "
    "independent ground truth for scoring the calls. Plate design details for the scifi-ATAC-seq "
    "libraries are in Table S1. Data accession gives the BioProject followed by the sequencing "
    "runs used here, and for the maize root library also the GEO sample. The synthetic benchmark "
    "is a titration series in which the genome of origin of every simulated read is known, so it "
    "provides exact truth rather than an estimate."
)


def build_rows(panels, wells, accessions, synth, root_accession):
    n_datasets, n_templates, alphas = synth
    alpha_labels = ", ".join("0" if a == 0 else f"{a / 100:.2f}" for a in alphas)
    synth_note = (
        f"{n_datasets} titration datasets (contamination fraction {alpha_labels}, "
        f"non-zero levels crossed with two contaminant genotypes)"
    )
    if n_templates:
        synth_note += f", {n_templates} template barcodes each"

    rows = []
    for key in ROW_ORDER:
        p = PROSE[key]
        if key.startswith("synthetic"):
            panel = ["B73", "Il14H", "Ki11"]
            n_wells = "Not applicable"
            mode = "Not applicable, benchmark of assignment and genotyping"
            material = p["material"] + ". " + synth_note
        else:
            panel = panels[key]
            n_wells = f"{wells[key]}" if wells.get(key) else "Not applicable"
            found = modes_on_disk(key)
            mode = " and ".join(found) if found else "Not applicable, genotyping only"
            material = p["material"]
        if p["accession"]:
            accession = p["accession"]
        elif key == "Root1":
            accession = ", ".join(root_accession)
        else:
            accession = ", ".join(accessions[key])
        rows.append([
            p["label"], p["assay"], material, len(panel), ", ".join(panel),
            n_wells, mode, p["role"], p["source"], p["pmid"], accession, p["figures"],
        ])
    return rows


def write_txt(path, rows):
    with open(path, "w", newline="") as fh:
        writer = csv.writer(fh, delimiter="\t", lineterminator="\n")
        writer.writerow(HEADER)
        writer.writerows(rows)
        fh.write("\n" + CAPTION + "\n")


def write_xlsx(path, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        print("openpyxl not available, skipping .xlsx", file=sys.stderr)
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "Table S2"
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="bottom")
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[0].font = Font(bold=True)
    widths = (26, 20, 46, 12, 34, 14, 20, 52, 18, 11, 34, 20)
    for column, width in zip("ABCDEFGHIJKL", widths):
        ws.column_dimensions[column].width = width
    ws.row_dimensions[1].height = 44
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 108

    caption_row = ws.max_row + 2
    ws.cell(row=caption_row, column=1, value=CAPTION).alignment = Alignment(
        wrap_text=True, vertical="top"
    )
    ws.merge_cells(
        start_row=caption_row, start_column=1, end_row=caption_row + 9, end_column=12
    )
    wb.save(path)
    return True


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--outdir", default=DEFAULT_OUTDIR)
    parser.add_argument(
        "--root-fastq", default=None,
        help="demultiplexed R1 FASTQ of the maize root library (Root1_rep1); when given and "
             "present, its first read header is checked against the cited SRA run",
    )
    args = parser.parse_args()

    panels = {k: read_panel(v) for k, v in PANEL_CONFIG.items()}
    wells = {k: count_wells(kind, path) for k, (kind, path) in WELL_SOURCE.items()}
    accessions = read_accessions()
    root_accession = read_root_runs(args.root_fastq)
    synth = read_synthetic()

    check(panels, wells, accessions, synth)
    print("self-check against the verified dataset inventory: PASS")

    rows = build_rows(panels, wells, accessions, synth, root_accession)

    os.makedirs(args.outdir, exist_ok=True)
    txt = os.path.join(args.outdir, "TableS2_dataset_inventory.txt")
    xlsx = os.path.join(args.outdir, "TableS2_dataset_inventory.xlsx")
    write_txt(txt, rows)
    print(f"wrote {txt}")
    if write_xlsx(xlsx, rows):
        print(f"wrote {xlsx}")

    width = max(len(r[0]) for r in rows)
    for row in rows:
        print(f"{row[0]:<{width}}  genomes={row[3]:>2}  wells={row[5]:<14}  "
              f"mode={row[6]:<40}  figs={row[11]}")

    pending = [(r[0], r[10]) for r in rows if r[10].startswith("[PENDING")]
    if pending:
        print("\n" + "=" * 78)
        print("PLACEHOLDERS REMAINING, these must be filled before submission:")
        for label, text in pending:
            print(f"  {label}: {text}")
        print("=" * 78)


if __name__ == "__main__":
    main()
